import html
import os
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from tempfile import gettempdir
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ARABIC_FONT = "Arial"


def _set_cell_shading(cell, fill="EAF2F8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text_direction(cell, rtl=True):
    tc_pr = cell._tc.get_or_add_tcPr()
    bidi = tc_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        tc_pr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")


def _set_paragraph_rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1")


def _set_run_font(run, size=11, bold=False):
    run.font.name = ARABIC_FONT
    run._element.rPr.rFonts.set(qn("w:cs"), ARABIC_FONT)
    run.font.size = Pt(size)
    run.bold = bold


def _add_rtl_paragraph(doc, text="", size=11, bold=False, space_after=4):
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    _set_run_font(run, size, bold)
    return p


def _add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    _set_run_font(run, 15 if level == 1 else 13, True)
    return p


def _safe_filename(text: str) -> str:
    text = re.sub(r"[\\/:*?\"<>|]+", "_", text)
    return text.strip() or "report"


def _item_time(item, tz):
    value = item.get("published_at") or item.get("discovered_at")
    if not value:
        return "غير متوفر"
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=ZoneInfo("UTC"))
        return dt.astimezone(tz).strftime("%Y-%m-%d %H:%M")
    except Exception:
        return str(value)


def _clean_content(value: str, max_chars=1200):
    if not value:
        return "لا يتوفر نص ملخص محفوظ لهذه المادة."
    text = re.sub(r"\s+", " ", value).strip()
    return text[:max_chars] + ("…" if len(text) > max_chars else "")


def _add_hyperlink(paragraph, text: str, url: str):
    if not url:
        run = paragraph.add_run(text)
        _set_run_font(run, 9, False)
        return
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), ARABIC_FONT)
    rFonts.set(qn("w:hAnsi"), ARABIC_FONT)
    rFonts.set(qn("w:cs"), ARABIC_FONT)
    rPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def create_daily_report(db, timezone_name="Asia/Aden", output_dir=None, report_title="مرصد المشهد الشرقي") -> str:
    tz = ZoneInfo(timezone_name)
    now = datetime.now(tz)
    items = db.get_today_items()
    sources = db.list_sources()
    stats = db.dashboard_stats()

    output_dir = output_dir or gettempdir()
    Path(output_dir).mkdir(parents=True, exist_ok=True)
    filename = f"موجز_أخبار_المهرة_{now.strftime('%Y-%m-%d_%H-%M')}.docx"
    path = str(Path(output_dir) / _safe_filename(filename))

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    styles["Normal"].font.name = ARABIC_FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), ARABIC_FONT)
    styles["Normal"].font.size = Pt(11)

    # Cover
    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run(report_title)
    _set_run_font(r, 24, True)

    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("الموجز الإخباري الشامل")
    _set_run_font(r, 19, True)

    p = doc.add_paragraph()
    _set_paragraph_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{now.strftime('%Y-%m-%d')} | {now.strftime('%H:%M')} بتوقيت عدن")
    _set_run_font(r, 12, False)

    _add_rtl_paragraph(doc, "", 10, False, 20)
    _add_rtl_paragraph(doc, "هذا التقرير مولد آليًا من قاعدة بيانات الرصد ويمكن فتحه وتحريره بالكامل في Microsoft Word.", 10, False, 10)
    doc.add_page_break()

    # Executive summary
    _add_heading(doc, "1. الملخص التنفيذي", 1)
    _add_rtl_paragraph(
        doc,
        f"حتى وقت إعداد التقرير، رصد النظام {len(items)} مادة إخبارية ضمن نطاق اليوم وفق توقيت {timezone_name}. "
        f"يضم التقرير المواد المرتبطة بمعايير الرصد، مع بيانات المصدر والوقت والرابط المتاح.",
        11,
    )

    _add_heading(doc, "2. مؤشرات الرصد", 1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["البيان", "القيمة"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        _set_cell_shading(cell)
        _set_cell_text_direction(cell)
        p = cell.paragraphs[0]
        _set_paragraph_rtl(p)
        run = p.add_run(text)
        _set_run_font(run, 11, True)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    metrics = [
        ("إجمالي المواد المرصودة اليوم", len(items)),
        ("المصادر المسجلة", stats.get("sources", 0)),
        ("المصادر النشطة", stats.get("enabled_sources", 0)),
        ("المصادر التي لديها أخطاء", stats.get("source_errors", 0)),
    ]
    for key, value in metrics:
        cells = table.add_row().cells
        for cell in cells:
            _set_cell_text_direction(cell)
        p = cells[0].paragraphs[0]
        _set_paragraph_rtl(p)
        _set_run_font(p.add_run(key), 10, False)
        p = cells[1].paragraphs[0]
        _set_paragraph_rtl(p)
        _set_run_font(p.add_run(str(value)), 10, True)

    # Sources status
    _add_heading(doc, "3. حالة المصادر", 1)
    for source in sources:
        status = "يعمل" if source.get("consecutive_errors", 0) == 0 else "متعثر"
        last = source.get("last_checked_at") or "لم يفحص بعد"
        text = f"{source.get('name', 'مصدر')} — الحالة: {status} — آخر فحص: {last}"
        if source.get("last_error"):
            text += f" — الخطأ: {source['last_error']}"
        _add_rtl_paragraph(doc, text, 10, False, 3)

    # News
    _add_heading(doc, "4. المواد الإخبارية المرصودة", 1)
    if not items:
        _add_rtl_paragraph(doc, "لم يتم رصد مواد إخبارية جديدة حتى وقت إعداد التقرير.", 11, False)
    else:
        source_counter = Counter(item.get("source_name", "مصدر غير معروف") for item in items)
        _add_rtl_paragraph(doc, "توزيع المواد حسب المصدر:", 11, True)
        for source_name, count in source_counter.most_common():
            _add_rtl_paragraph(doc, f"• {source_name}: {count} مادة", 10, False, 2)

        for idx, item in enumerate(items, 1):
            _add_heading(doc, f"{idx}. {item.get('title', 'بدون عنوان')}", 2)
            _add_rtl_paragraph(doc, f"المصدر: {item.get('source_name', 'غير معروف')}", 10, True, 2)
            _add_rtl_paragraph(doc, f"وقت النشر/الرصد: {_item_time(item, tz)}", 10, False, 2)
            score = item.get("mahra_score", 0)
            _add_rtl_paragraph(doc, f"درجة الصلة: {score}", 10, False, 2)
            _add_rtl_paragraph(doc, f"الملخص: {_clean_content(item.get('content', ''))}", 10, False, 4)
            link = item.get("link", "")
            p = doc.add_paragraph()
            _set_paragraph_rtl(p)
            run = p.add_run("رابط المادة الأصلية: ")
            _set_run_font(run, 10, True)
            if link:
                _add_hyperlink(p, link, link)
            doc.add_paragraph("────────────────────────────────────────")

    # Footer
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        _set_paragraph_rtl(footer)
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(f"مرصد المشهد الشرقي — تقرير مولد آليًا — {now.strftime('%Y-%m-%d %H:%M')}")
        _set_run_font(run, 8, False)

    doc.save(path)
    return path

